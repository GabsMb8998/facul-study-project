import json

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from content.models import Materia, Assunto
from exercises.models import Questao


class Command(BaseCommand):
    help = "Importa matérias, assuntos e questões de um arquivo Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "arquivo",
            type=str,
            help="Caminho do arquivo Excel (.xlsx)",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        caminho = options["arquivo"]

        try:
            workbook = load_workbook(
                caminho,
                data_only=True,
            )
        except FileNotFoundError:
            raise CommandError(
                f"Arquivo não encontrado: {caminho}"
            )
        except Exception as error:
            raise CommandError(
                f"Não foi possível abrir o Excel: {error}"
            )

        abas_obrigatorias = {
            "materias",
            "assuntos",
            "questoes",
        }

        abas_existentes = set(workbook.sheetnames)

        faltantes = abas_obrigatorias - abas_existentes

        if faltantes:
            raise CommandError(
                "Abas obrigatórias ausentes: "
                + ", ".join(sorted(faltantes))
            )

        try:
            materias = self.importar_materias(
                workbook["materias"]
            )

            assuntos = self.importar_assuntos(
                workbook["assuntos"],
                materias,
            )

            self.importar_questoes(
                workbook["questoes"],
                assuntos,
            )

        except Exception as error:
            raise CommandError(
                f"Erro durante a importação: {error}"
            )

        self.stdout.write(
            self.style.SUCCESS(
                "Importação concluída com sucesso!"
            )
        )

    def importar_materias(self, sheet):
        materias = {}

        for numero_linha, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            nome = self.texto(row[0])

            if not nome:
                continue

            descricao = self.texto(row[1])
            is_active = self.booleano(row[2])

            materia, _ = Materia.objects.update_or_create(
                nome=nome,
                defaults={
                    "descricao": descricao,
                    "is_active": is_active,
                },
            )

            materias[nome.lower()] = materia

        self.stdout.write(
            self.style.SUCCESS(
                f"Matérias importadas: {len(materias)}"
            )
        )

        return materias

    def importar_assuntos(self, sheet, materias):
        assuntos = {}

        for numero_linha, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            materia_nome = self.texto(row[0])
            nome = self.texto(row[1])

            if not materia_nome and not nome:
                continue

            if not materia_nome or not nome:
                raise CommandError(
                    f"Aba 'assuntos', linha {numero_linha}: "
                    "matéria e nome são obrigatórios."
                )

            materia = materias.get(
                materia_nome.lower()
            )

            if not materia:
                raise CommandError(
                    f"Aba 'assuntos', linha {numero_linha}: "
                    f"matéria '{materia_nome}' não encontrada."
                )

            conteudo = self.texto(row[2])
            ordem = self.numero(row[3])
            is_active = self.booleano(row[4])

            assunto, _ = Assunto.objects.update_or_create(
                materia=materia,
                nome=nome,
                defaults={
                    "conteudo": conteudo,
                    "ordem": ordem,
                    "is_active": is_active,
                },
            )

            chave = (
                materia_nome.lower(),
                nome.lower(),
            )

            assuntos[chave] = assunto

        self.stdout.write(
            self.style.SUCCESS(
                f"Assuntos importados: {len(assuntos)}"
            )
        )

        return assuntos

    def importar_questoes(self, sheet, assuntos):
        quantidade = 0

        for numero_linha, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            materia_nome = self.texto(row[0])
            assunto_nome = self.texto(row[1])
            enunciado = self.texto(row[2])

            if not materia_nome and not assunto_nome and not enunciado:
                continue

            if not materia_nome:
                raise CommandError(
                    f"Aba 'questoes', linha {numero_linha}: "
                    "matéria é obrigatória."
                )

            if not assunto_nome:
                raise CommandError(
                    f"Aba 'questoes', linha {numero_linha}: "
                    "assunto é obrigatório."
                )

            if not enunciado:
                raise CommandError(
                    f"Aba 'questoes', linha {numero_linha}: "
                    "enunciado é obrigatório."
                )

            tipo = self.texto(row[3])
            nivel = self.numero(row[4])
            vestibular = self.texto(row[5])
            dados = self.json(row[6])
            is_active = self.booleano(row[7])

            chave = (
                materia_nome.lower(),
                assunto_nome.lower(),
            )

            assunto = assuntos.get(chave)

            if not assunto:
                raise CommandError(
                    f"Aba 'questoes', linha {numero_linha}: "
                    f"assunto '{assunto_nome}' "
                    f"da matéria '{materia_nome}' "
                    "não encontrado."
                )

            Questao.objects.create(
                assunto=assunto,
                enunciado=enunciado,
                tipo=tipo,
                dados=dados,
                nivel=nivel,
                vestibular=vestibular,
                is_active=is_active,
            )

            quantidade += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Questões importadas: {quantidade}"
            )
        )

    @staticmethod
    def texto(valor):
        if valor is None:
            return ""

        return str(valor).strip()

    @staticmethod
    def booleano(valor):
        if isinstance(valor, bool):
            return valor

        if valor is None:
            return False

        return str(valor).strip().lower() in {
            "true",
            "1",
            "sim",
            "yes",
            "ativo",
            "active",
        }

    @staticmethod
    def numero(valor):
        if valor is None or valor == "":
            return None

        try:
            return int(valor)
        except (TypeError, ValueError):
            raise CommandError(
                f"Valor numérico inválido: {valor}"
            )

    @staticmethod
    def json(valor):
        if isinstance(valor, dict):
            return valor

        if not valor:
            raise CommandError(
                "O campo 'dados' não pode estar vazio."
            )

        try:
            return json.loads(str(valor))
        except json.JSONDecodeError as error:
            raise CommandError(
                f"JSON inválido no campo 'dados': {error}"
            )